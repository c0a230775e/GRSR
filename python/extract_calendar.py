import openpyxl
import json
import os

excel_path = os.path.join(os.path.dirname(__file__), "calendar.xlsx")
output_json = os.path.join(os.path.dirname(__file__), "../src/assets/timeSlots.json")
output_json = os.path.join(os.path.dirname(__file__), "../public/assets/timeSlots.json")

wb = openpyxl.load_workbook(excel_path)
ws = wb.active

results = []

# ★ 月は B1 のみを使う
month_value = ws.cell(row=1, column=2).value  # B1

# 数字だけ抽出（例：9月 → 9）
if month_value is not None:
    month_value = ''.join([c for c in str(month_value) if c.isdigit()])
else:
    month_value = None

# ★ カレンダー範囲：B2〜H7
for row in ws.iter_rows(min_row=2, max_row=7, min_col=2, max_col=8):
    for cell in row:
        if cell.value is None:
            continue

        lines = str(cell.value).splitlines()
        if len(lines) == 0:
            continue

        raw_date = lines[0].strip()
        dates = [d.strip() for d in raw_date.replace(" ", "").split("/") if d.strip()]

        labels = [line.strip() for line in lines[1:] if line.strip()]
        if len(labels) == 0:
            continue

        for date in dates:
            for label in labels:
                results.append({
                    "month": month_value,  # ← 全列に B1 の月を適用
                    "date": date,
                    "label": label
                })

with open(output_json, "w", encoding="utf-8") as f:
    json.dump(results, f, ensure_ascii=False, indent=2)

print("JSON_CREATED")
